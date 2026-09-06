"""Learn and render organization-native procurement output formats."""
from __future__ import annotations
import io
from typing import Any
from app.models import CommercialPosition
ROLE_KEYWORDS = {
 "executive_summary": ("executive","summary","overview","headline"),
 "supplier_position": ("supplier","ask","request","increase","position"),
 "economics": ("financial","cost","price","saving","spend","impact","econom"),
 "market": ("market","benchmark","index","external","commodity"),
 "supplier_performance": ("performance","otif","quality","service","kpi","srm"),
 "options": ("option","scenario","alternative","path"),
 "negotiation": ("negotiat","leverage","opening","target","walk-away","counter"),
 "risk": ("risk","concern","dependency","continuity","mitigation"),
 "decision": ("decision","approval","recommend","decision required"),
 "next_steps": ("next","action","timeline","follow-up","implementation"),
}
def _role_for(title: str, body: str) -> str:
 hay=f"{title} {body}".lower(); title_l=title.lower()
 if "negotiat" in title_l: return "negotiation"
 if "decision" in title_l or "approval" in title_l: return "decision"
 scores={r:sum(k in hay for k in ks) for r,ks in ROLE_KEYWORDS.items()}
 role=max(scores,key=scores.get); return role if scores[role] else "other"
def build_pptx_profile(file_bytes: bytes, filename: str)->dict[str,Any]:
 try:
  from pptx import Presentation
  prs=Presentation(io.BytesIO(file_bytes))
 except Exception as exc: raise ValueError("Could not read this PowerPoint format safely.") from exc
 slides=[]
 for idx,slide in enumerate(prs.slides,1):
  texts=[]
  for shape in slide.shapes:
   if getattr(shape,"has_text_frame",False):
    txt="\n".join(p.text for p in shape.text_frame.paragraphs).strip()
    if txt: texts.append(txt)
  if texts:
   title=texts[0][:300]; body="\n".join(texts[1:])[:1500]
   slides.append({"slide_number":idx,"title":title,"role":_role_for(title,body),"sample_text":body[:500]})
 if not slides: raise ValueError("No readable slide content was found in this PowerPoint.")
 return {"source_type":"pptx","source_filename":filename,"slide_count":len(prs.slides),"slides":slides}
def _alt_text(p):
 alt=getattr(p,"alternative_analysis",None)
 if not alt: return "No additional alternative path was resolved."
 items=getattr(alt,"alternatives",None) or []; out=[]
 for x in items[:5]:
  if isinstance(x,dict): out.append(f"- {x.get('name') or 'Alternative'}: {x.get('path') or ''}")
  else: out.append(f"- {x}")
 return "\n".join(out) or "No additional alternative path was resolved."
def _content_by_role(p: CommercialPosition)->dict[str,str]:
 fin=str(p.financial_impact) if p.financial_impact is not None else "Not safely calculable from supplied evidence"
 audit=getattr(p,"decision_audit",None); evidence=audit.evidence_integrity_status if audit else "UNKNOWN"
 cockpit=getattr(p,"decision_cockpit",None) or {}; dims=getattr(getattr(p,"negotiation_playbook",None),"dimensions",None) or []
 negotiation=(f"Opening position: {p.opening_position or 'Not stated.'}\n" + "\n".join(f"{d.get('dimension')}: target={d.get('target')}; walk-away={d.get('walk_away')}" for d in dims[:8] if isinstance(d,dict))).strip()
 blockers=getattr(getattr(p,"control_tower",None),"critical_before_action",None) or []
 return {"executive_summary":f"Recommendation: {p.recommendation}\nConfidence: {p.confidence.level}\nEvidence integrity: {evidence}","supplier_position":p.opening_position or "Supplier/request position is contained in the validated evidence.","economics":f"Financial impact: {fin}","market":str((cockpit.get("economics") or {}).get("headline") or "Market evidence: see the validated evidence ledger."),"supplier_performance":"Supplier performance: use validated supplier evidence; no unsupported KPI is introduced.","options":_alt_text(p),"negotiation":negotiation,"risk":"\n".join(f"- {x}" for x in blockers[:8]) or "No additional pre-action blocker recorded.","decision":f"Decision required: {p.recommendation}\nDecision changer: {p.disconfirming_condition}","next_steps":str(cockpit.get("next_move") or "Review evidence and complete the approval-gated next action."),"other":"Use the underlying VendorEdge decision for the complete auditable detail."}
def render_organisation_format(p: CommercialPosition, profile: dict[str,Any])->dict[str,Any]:
 contents=_content_by_role(p); slides=[]
 for slide in profile.get("slides",[]):
  role=slide.get("role","other"); slides.append({"slide_number":slide.get("slide_number"),"title":slide.get("title") or "Untitled","role":role,"content":contents.get(role,contents["other"])})
 return {"format":"organisation_native","title":profile.get("source_filename","Organization format"),"slides":slides,"method":"Deterministic rendering into the organization's learned structure; no new facts, calculations or LLM call."}
