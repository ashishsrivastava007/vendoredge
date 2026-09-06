import io
from types import SimpleNamespace
from app.pipeline.file_extraction import extract_text_from_pptx
from app.pipeline.organisation_formats import build_pptx_profile, render_organisation_format


def make_pptx():
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    for title, body in [
        ("Negotiation Strategy", "Supplier ask and opening position"),
        ("Financial Impact", "Savings and price impact"),
        ("Decision Required", "Management approval")
    ]:
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = title
        slide.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(2)).text_frame.text = body
    out = io.BytesIO(); prs.save(out); return out.getvalue()


def position():
    return SimpleNamespace(
        recommendation="Hold price pending evidence", confidence=SimpleNamespace(level="high"),
        financial_impact="USD 100,000", opening_position="Reject unsupported increase",
        disconfirming_condition="Verified cost evidence supports increase",
        decision_audit=SimpleNamespace(evidence_integrity_status="CERTIFIED"),
        decision_cockpit={"next_move":"Request cost evidence", "economics":{"headline":"USD 100,000 exposure"}},
        decision_passport={},
        negotiation_playbook=SimpleNamespace(dimensions=[{"dimension":"Price","target":"0%","walk_away":">3%"}]),
        control_tower=SimpleNamespace(critical_before_action=["Validate cost evidence"]),
        alternative_analysis=None,
    )


def test_pptx_extraction_and_profile_are_deterministic():
    raw = make_pptx()
    text = extract_text_from_pptx(raw)
    assert "Negotiation Strategy" in text and "Financial Impact" in text
    profile = build_pptx_profile(raw, "Negotiation Deck.pptx")
    assert profile["slide_count"] == 3
    assert [s["role"] for s in profile["slides"]] == ["negotiation", "economics", "decision"]


def test_organisation_render_preserves_slide_structure_and_uses_decision():
    profile = {"source_filename":"Negotiation Deck.pptx", "slides":[
        {"slide_number":1,"title":"Negotiation Strategy","role":"negotiation"},
        {"slide_number":2,"title":"Financial Impact","role":"economics"},
    ]}
    out = render_organisation_format(position(), profile)
    assert [x["title"] for x in out["slides"]] == ["Negotiation Strategy", "Financial Impact"]
    assert "Reject unsupported increase" in out["slides"][0]["content"]
    assert "USD 100,000" in out["slides"][1]["content"]
    assert "no new facts" in out["method"]
